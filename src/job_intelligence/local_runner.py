from __future__ import annotations

import argparse
import contextlib
import hashlib
import json
import os
import re
import shutil
import signal
import socket
import sqlite3
import subprocess
import sys
import time
import zipfile
from datetime import UTC, datetime, timedelta
from pathlib import Path
from typing import IO, Any, Iterator
from urllib.parse import urlsplit

from openpyxl import load_workbook

from .excel_export import REQUIRED_SHEETS
from .source_config import load_source_config

_RUN_PATTERN = re.compile(r"^run-\d{8}-\d{6}-attempt-\d+$")
_TRANSIENT = re.compile(
    r"(?i)(timeout|timed out|connection|dns|temporar|429|too many requests|"
    r"502|503|504|name resolution|network is unreachable)"
)
_SECRET_PATTERNS = (
    re.compile(r"gh[pousr]_[A-Za-z0-9_]{20,}"),
    re.compile(r"AKIA[0-9A-Z]{16}"),
    re.compile(r"(?i)(access[_-]?token|client[_-]?secret|password)\s*[:=]\s*\S+"),
    re.compile(r"(?i)\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b"),
)


def now_utc() -> datetime:
    return datetime.now(UTC)


def redact(value: str) -> str:
    result = value
    for pattern in _SECRET_PATTERNS:
        result = pattern.sub("[REDACTED]", result)
    return result


def is_transient_failure(value: str) -> bool:
    return bool(_TRANSIENT.search(value))


def contained_child(root: Path, child: Path) -> Path:
    resolved_root = root.resolve()
    resolved_child = child.resolve()
    if resolved_child == resolved_root or resolved_root not in resolved_child.parents:
        raise ValueError(f"output must be a child of data root: {resolved_child}")
    if resolved_root == Path(resolved_root.anchor):
        raise ValueError("data root must not be a drive or filesystem root")
    return resolved_child


def run_name(stamp: datetime, attempt: int) -> str:
    if attempt < 1:
        raise ValueError("attempt must be >= 1")
    return f"run-{stamp.astimezone(UTC):%Y%m%d-%H%M%S}-attempt-{attempt}"


def _process_exists(pid: int) -> bool:
    if pid <= 0:
        return False
    try:
        os.kill(pid, 0)
    except OSError:
        return False
    return True


@contextlib.contextmanager
def exclusive_lock(path: Path) -> Iterator[IO[bytes]]:
    path.parent.mkdir(parents=True, exist_ok=True)
    handle = path.open("a+b")
    handle.seek(0, os.SEEK_END)
    if handle.tell() == 0:
        handle.write(b"0")
        handle.flush()
    handle.seek(0)
    try:
        if os.name == "nt":
            import msvcrt

            msvcrt.locking(handle.fileno(), msvcrt.LK_NBLCK, 1)
        else:
            import fcntl

            fcntl.flock(handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
    except OSError as exc:
        handle.close()
        raise RuntimeError(f"another local public run holds {path}") from exc
    try:
        handle.seek(0)
        handle.truncate()
        handle.write(
            json.dumps(
                {"pid": os.getpid(), "host": socket.gethostname(), "started": now_utc().isoformat()}
            ).encode("utf-8")
        )
        handle.flush()
        yield handle
    finally:
        handle.seek(0)
        handle.truncate()
        if os.name == "nt":
            import msvcrt

            handle.seek(0)
            msvcrt.locking(handle.fileno(), msvcrt.LK_UNLCK, 1)
        else:
            import fcntl

            fcntl.flock(handle.fileno(), fcntl.LOCK_UN)
        handle.close()


def _kill_tree(process: subprocess.Popen[str]) -> None:
    if process.poll() is not None:
        return
    if os.name == "nt":
        subprocess.run(
            ["taskkill", "/PID", str(process.pid), "/T", "/F"],
            capture_output=True,
            text=True,
            check=False,
        )
    else:
        os.killpg(process.pid, signal.SIGKILL)
    try:
        process.wait(timeout=15)
    except subprocess.TimeoutExpired:
        process.kill()


def run_logged(
    command: list[str],
    *,
    cwd: Path,
    log: IO[str],
    timeout_seconds: float,
) -> int:
    if timeout_seconds <= 0:
        raise TimeoutError("overall timeout expired before child start")
    flags = subprocess.CREATE_NEW_PROCESS_GROUP if os.name == "nt" else 0
    process = subprocess.Popen(
        command,
        cwd=cwd,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        encoding="utf-8",
        errors="replace",
        creationflags=flags,
        start_new_session=os.name != "nt",
    )
    try:
        output, _ = process.communicate(timeout=timeout_seconds)
    except subprocess.TimeoutExpired as exc:
        _kill_tree(process)
        raise TimeoutError(f"child exceeded timeout: {command[1]}") from exc
    except BaseException:
        _kill_tree(process)
        raise
    for line in output.splitlines():
        log.write(f"{now_utc().isoformat()} {redact(line)}\n")
    log.flush()
    return int(process.returncode)


def _source_hosts(config_path: Path) -> set[str]:
    config = load_source_config(config_path)
    hosts: set[str] = set()
    for source in config.sources:
        if not source.enabled:
            continue
        for key in ("base_url", "url", "api_url", "listing_url", "page_url_template"):
            raw = source.options.get(key)
            if isinstance(raw, str):
                host = urlsplit(raw).hostname
                if host:
                    hosts.add(host)
    return hosts


def connectivity_preflight(config_path: Path) -> list[str]:
    failures: list[str] = []
    for host in sorted(_source_hosts(config_path)):
        try:
            socket.getaddrinfo(host, 443, type=socket.SOCK_STREAM)
        except OSError as exc:
            failures.append(f"{host}: {type(exc).__name__}")
    return failures


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def verify_output(output: Path, started: datetime) -> dict[str, Any]:
    required = {
        "status": output / "status.json",
        "summary": output / "SUMMARY.md",
        "log": output / "run.log",
        "database": output / "jobs.db",
        "workbook": output / "Piping_E3D_Jobs.xlsx",
        "bundle": output / "Piping_E3D_Daily_Bundle.zip",
    }
    errors = [f"missing or empty {name}" for name, path in required.items() if not path.is_file() or path.stat().st_size <= 0]
    if errors:
        return {"result": "FAIL", "errors": errors}
    status = json.loads(required["status"].read_text(encoding="utf-8"))
    connection = sqlite3.connect(f"file:{required['database']}?mode=ro", uri=True)
    try:
        integrity = connection.execute("PRAGMA integrity_check").fetchall()
        jobs = connection.execute("SELECT COUNT(*) FROM jobs").fetchone()[0]
        private_count = connection.execute("SELECT COUNT(*) FROM private_imports").fetchone()[0]
        gmail_messages = connection.execute("SELECT COUNT(*) FROM gmail_alert_messages").fetchone()[0]
        gmail_jobs = connection.execute("SELECT COUNT(*) FROM gmail_alert_jobs").fetchone()[0]
        evidence = connection.execute(
            "SELECT file_path, sha256, size_bytes FROM source_evidence"
        ).fetchall()
    finally:
        connection.close()
    if integrity != [("ok",)]:
        errors.append("SQLite integrity_check failed")
    if private_count or gmail_messages or gmail_jobs:
        errors.append("private or Gmail tables are not empty")
    bad_evidence = 0
    for raw_path, expected_hash, expected_size in evidence:
        path = Path(raw_path)
        if not path.is_absolute():
            path = output / path
        if not path.is_file() or path.stat().st_size != expected_size or _sha256(path).casefold() != str(expected_hash).casefold():
            bad_evidence += 1
    if bad_evidence:
        errors.append(f"{bad_evidence} evidence files failed hash/size verification")
    workbook = load_workbook(required["workbook"], read_only=True, data_only=True)
    try:
        missing_sheets = [sheet for sheet in REQUIRED_SHEETS if sheet not in workbook.sheetnames]
        sheet_count = len(workbook.sheetnames)
    finally:
        workbook.close()
    if missing_sheets:
        errors.append("missing workbook sheets: " + ", ".join(missing_sheets))
    with zipfile.ZipFile(required["bundle"]) as archive:
        bad_member = archive.testzip()
        zip_names = archive.namelist()
    if bad_member:
        errors.append(f"corrupt ZIP member: {bad_member}")
    prohibited = re.compile(r"(?i)(credentials|token|private-output|gmail-evidence|\.env$|\.venv)")
    if any(prohibited.search(name) for name in zip_names):
        errors.append("ZIP contains prohibited private/credential path")
    stale = [
        str(path.relative_to(output))
        for path in output.rglob("*")
        if path.is_file()
        and datetime.fromtimestamp(path.stat().st_mtime, UTC) < started
    ]
    if stale:
        errors.append(f"stale files detected: {len(stale)}")
    if status.get("exit_code") != 0 or status.get("verified") is not True:
        errors.append("production status is not verified PASS")
    manifest_lines = []
    for path in sorted(p for p in output.rglob("*") if p.is_file() and p.name not in {"artifact-manifest.sha256", "local-runner-status.json"}):
        manifest_lines.append(f"{_sha256(path)}  {path.relative_to(output)}")
    (output / "artifact-manifest.sha256").write_text("\n".join(manifest_lines) + "\n", encoding="utf-8")
    return {
        "result": "FAIL" if errors else "PASS",
        "errors": errors,
        "jobs": jobs,
        "evidence_files": len(evidence),
        "sheet_count": sheet_count,
        "required_sheet_count": len(REQUIRED_SHEETS),
        "zip_entries": len(zip_names),
        "manifest_entries": len(manifest_lines),
    }


def retention_cleanup(data_root: Path, *, keep_days: int, keep_runs: int) -> list[str]:
    if keep_days < 1 or keep_runs < 1:
        raise ValueError("retention bounds must be positive")
    root = data_root.resolve()
    candidates = sorted(
        (path for path in root.iterdir() if path.is_dir() and _RUN_PATTERN.fullmatch(path.name)),
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )
    cutoff = now_utc() - timedelta(days=keep_days)
    removed: list[str] = []
    for index, path in enumerate(candidates):
        contained_child(root, path)
        modified = datetime.fromtimestamp(path.stat().st_mtime, UTC)
        if index >= keep_runs and modified < cutoff:
            shutil.rmtree(path)
            removed.append(path.name)
    return removed


def _git(repo: Path, *args: str) -> str:
    return subprocess.check_output(["git", *args], cwd=repo, text=True).strip()


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Run the public-only Piping/E3D daily pipeline safely on Windows.",
        epilog=(
            "Emergency stop: Stop-Process -Id <wrapper-child-pid>. The runner terminates "
            "its process tree on timeout. Never delete the repository; outputs are isolated."
        ),
    )
    parser.add_argument("--repo", required=True)
    parser.add_argument("--data-root", required=True)
    parser.add_argument("--sources", default="config/sources.yaml")
    parser.add_argument("--timeout-minutes", type=int, default=120)
    parser.add_argument("--retries", type=int, default=1)
    parser.add_argument("--min-free-gib", type=float, default=5.0)
    parser.add_argument("--keep-days", type=int, default=14)
    parser.add_argument("--keep-runs", type=int, default=14)
    parser.add_argument("--allow-no-sources", action="store_true")
    parser.add_argument("--skip-connectivity", action="store_true")
    return parser


def main() -> int:
    args = build_parser().parse_args()
    repo = Path(args.repo).resolve()
    data_root = Path(args.data_root).resolve()
    if sys.version_info[:2] != (3, 11):
        raise SystemExit("Python 3.11 is required")
    if not (repo / ".git").exists() or not (repo / "scripts/run_cloud_daily.py").is_file():
        raise SystemExit("invalid repository path")
    if Path(sys.executable).resolve() != (repo / ".venv/Scripts/python.exe").resolve():
        raise SystemExit("runner must use the repository-local .venv Python")
    data_root.mkdir(parents=True, exist_ok=True)
    contained_child(data_root.parent, data_root)
    free = shutil.disk_usage(data_root).free
    if free < args.min_free_gib * 1024**3:
        raise SystemExit("minimum free-space gate failed")
    branch = _git(repo, "branch", "--show-current")
    commit = _git(repo, "rev-parse", "HEAD")
    if branch != "local/windows-runner-setup":
        raise SystemExit(f"unexpected branch: {branch}")
    if _git(repo, "status", "--porcelain"):
        raise SystemExit("tracked or untracked repository changes detected")
    config = (repo / args.sources).resolve() if not Path(args.sources).is_absolute() else Path(args.sources).resolve()
    if repo not in config.parents:
        raise SystemExit("sources configuration must be inside repository")
    source_config = load_source_config(config)
    if any(source.source_type.startswith("gmail") for source in source_config.sources if source.enabled):
        raise SystemExit("Gmail/private source is not permitted in the public runner")
    if not args.skip_connectivity:
        failures = connectivity_preflight(config)
        if failures:
            raise SystemExit("connectivity preflight failed: " + "; ".join(failures))
    stamp = now_utc()
    deadline = time.monotonic() + args.timeout_minutes * 60
    lock_path = data_root / ".public-daily.lock"
    final_status: dict[str, Any] = {"result": "FAIL", "commit": commit, "branch": branch}
    with exclusive_lock(lock_path):
        for attempt in range(1, args.retries + 2):
            output = contained_child(data_root, data_root / run_name(stamp, attempt))
            if output.exists():
                raise SystemExit(f"refusing to overwrite existing run: {output}")
            output.mkdir(parents=True)
            log_path = output / "local-runner.log"
            started = now_utc()
            with log_path.open("w", encoding="utf-8") as log:
                log.write(f"{started.isoformat()} START commit={commit} attempt={attempt}\n")
                command = [sys.executable, str(repo / "scripts/run_cloud_daily.py"), "--sources", str(config), "--output", str(output)]
                if args.allow_no_sources:
                    command.append("--allow-no-sources")
                try:
                    code = run_logged(command, cwd=repo, log=log, timeout_seconds=deadline - time.monotonic())
                    coverage_code = run_logged(
                        [sys.executable, str(repo / "scripts/append_daily_coverage.py"), "--output", str(output), "--config-dir", str(repo / "config")],
                        cwd=repo,
                        log=log,
                        timeout_seconds=deadline - time.monotonic(),
                    )
                    proof = verify_output(output, started)
                    final_status = {**proof, "daily_exit": code, "coverage_exit": coverage_code, "attempt": attempt, "commit": commit, "branch": branch, "started_at": started.isoformat(), "ended_at": now_utc().isoformat(), "output": str(output)}
                except Exception as exc:
                    final_status = {"result": "FAIL", "error": redact(f"{type(exc).__name__}: {exc}"), "attempt": attempt, "commit": commit, "branch": branch, "output": str(output)}
                (output / "local-runner-status.json").write_text(json.dumps(final_status, indent=2, sort_keys=True) + "\n", encoding="utf-8")
            if final_status.get("result") == "PASS" and final_status.get("daily_exit") == 0 and final_status.get("coverage_exit") == 0:
                retention_cleanup(data_root, keep_days=args.keep_days, keep_runs=args.keep_runs)
                print(json.dumps(final_status, sort_keys=True))
                return 0
            status_text = json.dumps(final_status, sort_keys=True)
            if attempt > args.retries or not is_transient_failure(status_text):
                print(status_text)
                return 1
            delay = min(60 * (2 ** (attempt - 1)), 300)
            if time.monotonic() + delay >= deadline:
                print(status_text)
                return 1
            time.sleep(delay)
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
