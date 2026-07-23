from __future__ import annotations

from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font

from .database import connect
from .private_import import init_private_import_tables
from .profile_enrichment import (
    fetch_enriched_jobs,
    refresh_job_enrichment,
)

REQUIRED_SHEETS = (
    "New_Today",
    "High_Priority",
    "Worldwide_Active",
    "All_Active",
    "Cross_Source_Duplicates",
    "Global_Summary",
    "Manual_Review",
    "Applied",
    "Follow_Up",
    "Expired",
    "Recruiter_Contacts",
    "Source_Health",
    "Source_Evidence",
    "Private_Imports",
    "Run_Proof",
)
_FORMULA_PREFIXES = ("=", "+", "-", "@")


_WORLDWIDE_COLUMNS = (
    "job_key",
    "title",
    "normalized_role",
    "company",
    "location",
    "city",
    "country",
    "apply_url",
    "source_name",
    "source_names",
    "source_type",
    "published_at",
    "closing_date",
    "experience_required",
    "software",
    "sector",
    "employment_mode",
    "salary_text",
    "match_score",
    "match_reasons",
    "gaps",
    "contact_confidence",
    "recruiter_email",
    "duplicate_status",
    "duplicate_source_count",
    "canonical_job_key",
    "application_status",
    "found_at",
    "last_seen_at",
)


def _empty_jobs_frame() -> pd.DataFrame:
    return pd.DataFrame(columns=_WORLDWIDE_COLUMNS)


def _safe_excel_value(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    inspected = value.lstrip(" \t\r\n")
    if inspected.startswith(_FORMULA_PREFIXES):
        return "'" + value
    return value


def _safe_excel_frame(frame: pd.DataFrame) -> pd.DataFrame:
    safe = frame.copy()
    for column in safe.columns:
        safe[column] = safe[column].map(_safe_excel_value)
    return safe


def _format_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for column_cells in sheet.columns:
            values = [str(cell.value or "") for cell in list(column_cells)[:200]]
            maximum = max((len(value) for value in values), default=0)
            width = min(max(maximum + 2, 10), 55)
            sheet.column_dimensions[column_cells[0].column_letter].width = width
    workbook.save(path)


def _worldwide_view(jobs: pd.DataFrame) -> pd.DataFrame:
    for column in _WORLDWIDE_COLUMNS:
        if column not in jobs.columns:
            jobs[column] = ""
    return jobs.loc[:, list(_WORLDWIDE_COLUMNS)]


def _global_summary(active: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, object]] = [
        {"dimension": "TOTAL", "value": "active jobs", "job_count": len(active)}
    ]
    for column, dimension in (
        ("country", "COUNTRY"),
        ("normalized_role", "NORMALIZED_ROLE"),
        ("sector", "SECTOR"),
        ("source_name", "SOURCE"),
        ("priority", "PRIORITY"),
    ):
        if column not in active.columns:
            continue
        values = active[column].fillna("").astype(str).str.strip()
        counts = values[values.ne("")].value_counts()
        rows.extend(
            {
                "dimension": dimension,
                "value": value,
                "job_count": int(count),
            }
            for value, count in counts.items()
        )
    return pd.DataFrame(rows, columns=("dimension", "value", "job_count"))


def export_excel(db_path: str | Path, output_path: str | Path) -> Path:
    init_private_import_tables(db_path)
    refresh_job_enrichment(db_path)
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    jobs = pd.DataFrame(fetch_enriched_jobs(db_path))
    if jobs.empty:
        jobs = _empty_jobs_frame()

    with connect(db_path) as connection:
        source_health = pd.read_sql_query(
            "SELECT * FROM source_health ORDER BY last_attempt_at DESC",
            connection,
        )
        source_evidence = pd.read_sql_query(
            "SELECT * FROM source_evidence ORDER BY fetched_at DESC",
            connection,
        )
        private_imports = pd.read_sql_query(
            "SELECT * FROM private_imports ORDER BY imported_at DESC",
            connection,
        )
        run_proof = pd.read_sql_query(
            "SELECT * FROM runs ORDER BY started_at DESC",
            connection,
        )

    review_job_keys = set(
        private_imports.loc[
            private_imports.get("review_required", pd.Series(dtype="int64")).eq(1),
            "job_key",
        ].astype(str)
    )
    today = datetime.now(UTC).date().isoformat()
    found_dates = jobs.get("found_at", pd.Series(dtype="string")).astype(str)
    new_today = jobs[found_dates.str.startswith(today, na=False)]
    priorities = jobs.get("priority", pd.Series(dtype="string"))
    high_priority = jobs[priorities.isin(["critical", "high"])]
    statuses = jobs.get("application_status", pd.Series(dtype="string"))
    active = jobs[~statuses.isin(["expired", "rejected"])]
    confidence = jobs.get("contact_confidence", pd.Series(dtype="string"))
    job_keys = jobs.get("job_key", pd.Series(dtype="string")).astype(str)
    manual_review = jobs[
        confidence.isin(["PUBLIC_UNVERIFIED", "PATTERN_SUGGESTION"])
        | statuses.eq("review_required")
        | job_keys.isin(review_job_keys)
    ]
    applied = jobs[statuses.eq("applied")]
    follow_up = jobs[statuses.eq("follow_up")]
    expired = jobs[statuses.eq("expired")]
    emails = jobs.get("recruiter_email", pd.Series(dtype="string"))
    contacts = jobs[emails.fillna("").astype(str).str.len() > 0]
    duplicates = jobs[
        jobs.get("duplicate_status", pd.Series(dtype="string")).eq("cross_source")
    ]
    worldwide_active = _worldwide_view(active.copy())

    sheets = {
        "New_Today": _worldwide_view(new_today.copy()),
        "High_Priority": _worldwide_view(high_priority.copy()),
        "Worldwide_Active": worldwide_active,
        "All_Active": active,
        "Cross_Source_Duplicates": _worldwide_view(duplicates.copy()),
        "Global_Summary": _global_summary(active),
        "Manual_Review": _worldwide_view(manual_review.copy()),
        "Applied": _worldwide_view(applied.copy()),
        "Follow_Up": _worldwide_view(follow_up.copy()),
        "Expired": _worldwide_view(expired.copy()),
        "Recruiter_Contacts": _worldwide_view(contacts.copy()),
        "Source_Health": source_health,
        "Source_Evidence": source_evidence,
        "Private_Imports": private_imports,
        "Run_Proof": run_proof,
    }
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for sheet_name, frame in sheets.items():
            _safe_excel_frame(frame).to_excel(
                writer,
                sheet_name=sheet_name,
                index=False,
            )

    _format_workbook(output)
    return output


def verify_excel(path: str | Path) -> list[str]:
    workbook_path = Path(path)
    errors: list[str] = []
    if not workbook_path.exists():
        return [f"Excel output does not exist: {workbook_path}"]
    if workbook_path.stat().st_size == 0:
        return [f"Excel output is empty: {workbook_path}"]

    workbook = load_workbook(workbook_path, read_only=True)
    missing = [sheet for sheet in REQUIRED_SHEETS if sheet not in workbook.sheetnames]
    if missing:
        errors.append(f"Missing sheets: {', '.join(missing)}")
    if "Worldwide_Active" in workbook.sheetnames:
        headers = [cell.value for cell in next(workbook["Worldwide_Active"].iter_rows())]
        missing_columns = [
            column for column in _WORLDWIDE_COLUMNS if column not in headers
        ]
        if missing_columns:
            errors.append(
                "Worldwide_Active missing columns: " + ", ".join(missing_columns)
            )
    return errors
