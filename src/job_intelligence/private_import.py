from __future__ import annotations

import hashlib
import mimetypes
import re
import shutil
import uuid
from dataclasses import dataclass, field
from email import policy
from email.parser import BytesParser
from pathlib import Path
from typing import Iterable

from bs4 import BeautifulSoup
from docx import Document
from PIL import Image
from pypdf import PdfReader

from .database import connect, upsert_job
from .manual_import import create_manual_job
from .models import utc_now_iso
from .proof import init_proof_tables

SUPPORTED_EXTENSIONS = {
    ".txt",
    ".md",
    ".csv",
    ".pdf",
    ".docx",
    ".eml",
    ".png",
    ".jpg",
    ".jpeg",
    ".webp",
    ".tif",
    ".tiff",
    ".bmp",
}
_IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff", ".bmp"}
_URL_PATTERN = re.compile(r"https?://[^\s<>\]\[\"']+", re.IGNORECASE)
_LABEL_PATTERNS = {
    "company": re.compile(r"^(?:company|organisation|organization|employer)\s*:\s*(.+)$", re.I),
    "location": re.compile(r"^(?:location|job location|work location)\s*:\s*(.+)$", re.I),
    "title": re.compile(r"^(?:job title|position|role)\s*:\s*(.+)$", re.I),
}
_JOB_WORDS = (
    "piping",
    "pipeline",
    "e3d",
    "pdms",
    "layout",
    "mechanical",
    "designer",
    "engineer",
    "lead",
    "manager",
    "supervisor",
)


@dataclass(frozen=True, slots=True)
class ExtractedDocument:
    text: str
    method: str
    mime_type: str
    warnings: tuple[str, ...] = ()


@dataclass(frozen=True, slots=True)
class PrivateImportResult:
    input_path: str
    status: str
    job_key: str = ""
    created: bool = False
    sha256: str = ""
    stored_path: str = ""
    review_required: bool = False
    warnings: tuple[str, ...] = ()
    error_message: str = ""


@dataclass(slots=True)
class FolderImportSummary:
    attempted: int = 0
    created: int = 0
    updated: int = 0
    duplicates: int = 0
    failed: int = 0
    review_required: int = 0
    results: list[PrivateImportResult] = field(default_factory=list)


def init_private_import_tables(db_path: str | Path) -> None:
    init_proof_tables(db_path)
    with connect(db_path) as connection:
        connection.executescript(
            """
            CREATE TABLE IF NOT EXISTS private_imports (
                import_id TEXT PRIMARY KEY,
                sha256 TEXT NOT NULL UNIQUE,
                original_path TEXT NOT NULL,
                original_name TEXT NOT NULL,
                stored_path TEXT NOT NULL,
                mime_type TEXT NOT NULL DEFAULT '',
                size_bytes INTEGER NOT NULL,
                extraction_method TEXT NOT NULL,
                extracted_chars INTEGER NOT NULL DEFAULT 0,
                review_required INTEGER NOT NULL DEFAULT 0,
                warnings TEXT NOT NULL DEFAULT '',
                job_key TEXT NOT NULL DEFAULT '',
                status TEXT NOT NULL,
                imported_at TEXT NOT NULL
            );

            CREATE INDEX IF NOT EXISTS idx_private_imports_job
                ON private_imports(job_key);
            CREATE INDEX IF NOT EXISTS idx_private_imports_status
                ON private_imports(status, imported_at);
            """
        )


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _extract_pdf(path: Path) -> ExtractedDocument:
    reader = PdfReader(str(path))
    pages = [(page.extract_text() or "").strip() for page in reader.pages]
    text = "\n\n".join(page for page in pages if page)
    warnings: list[str] = []
    if len(text.strip()) < 80:
        warnings.append(
            "PDF contains little extractable text; scanned pages require OCR or manual review"
        )
    return ExtractedDocument(
        text=text,
        method="pypdf",
        mime_type="application/pdf",
        warnings=tuple(warnings),
    )


def _extract_docx(path: Path) -> ExtractedDocument:
    document = Document(str(path))
    blocks: list[str] = [paragraph.text.strip() for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            blocks.append(" | ".join(cell.text.strip() for cell in row.cells))
    text = "\n".join(block for block in blocks if block)
    return ExtractedDocument(
        text=text,
        method="python-docx",
        mime_type=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
    )


def _extract_eml(path: Path) -> ExtractedDocument:
    message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
    blocks: list[str] = []
    for label, value in (
        ("Subject", message.get("subject", "")),
        ("From", message.get("from", "")),
        ("To", message.get("to", "")),
        ("Date", message.get("date", "")),
    ):
        if value:
            blocks.append(f"{label}: {value}")

    warnings: list[str] = []
    body_added = False
    if message.is_multipart():
        for part in message.walk():
            disposition = part.get_content_disposition()
            content_type = part.get_content_type()
            if disposition == "attachment":
                filename = part.get_filename() or "unnamed attachment"
                warnings.append(f"Email attachment not imported automatically: {filename}")
                continue
            if content_type == "text/plain" and not body_added:
                blocks.append(str(part.get_content()))
                body_added = True
        if not body_added:
            for part in message.walk():
                if part.get_content_type() == "text/html":
                    html = str(part.get_content())
                    blocks.append(BeautifulSoup(html, "html.parser").get_text("\n"))
                    body_added = True
                    break
    else:
        content = str(message.get_content())
        if message.get_content_type() == "text/html":
            content = BeautifulSoup(content, "html.parser").get_text("\n")
        blocks.append(content)

    return ExtractedDocument(
        text="\n".join(block.strip() for block in blocks if block and block.strip()),
        method="email-parser",
        mime_type="message/rfc822",
        warnings=tuple(warnings),
    )


def _extract_image(
    path: Path,
    *,
    ocr: bool,
    tesseract_cmd: str | None,
) -> ExtractedDocument:
    if not ocr:
        raise ValueError("image import requires --ocr")
    try:
        import pytesseract
    except ImportError as exc:
        raise RuntimeError("pytesseract is not installed") from exc
    if tesseract_cmd:
        pytesseract.pytesseract.tesseract_cmd = tesseract_cmd
    try:
        with Image.open(path) as image:
            text = pytesseract.image_to_string(image.convert("RGB"))
    except Exception as exc:
        raise RuntimeError(
            "OCR failed; verify that Tesseract OCR is installed and accessible"
        ) from exc
    warnings = (
        "OCR text requires manual review before relying on inferred vacancy fields",
    )
    return ExtractedDocument(
        text=text.strip(),
        method="tesseract-ocr",
        mime_type=Image.MIME.get(Image.open(path).format, "image/*"),
        warnings=warnings,
    )


def extract_document(
    path: str | Path,
    *,
    ocr: bool = False,
    tesseract_cmd: str | None = None,
    max_bytes: int = 25_000_000,
) -> ExtractedDocument:
    source = Path(path)
    if not source.is_file():
        raise FileNotFoundError(f"input file does not exist: {source}")
    if source.stat().st_size > max_bytes:
        raise ValueError(f"input file exceeds {max_bytes} bytes: {source}")
    extension = source.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"unsupported private import type: {extension or '<none>'}")

    if extension in {".txt", ".md", ".csv"}:
        mime_type = mimetypes.guess_type(source.name)[0] or "text/plain"
        return ExtractedDocument(
            text=_decode_text(source.read_bytes()).strip(),
            method="text-decoder",
            mime_type=mime_type,
        )
    if extension == ".pdf":
        return _extract_pdf(source)
    if extension == ".docx":
        return _extract_docx(source)
    if extension == ".eml":
        return _extract_eml(source)
    return _extract_image(source, ocr=ocr, tesseract_cmd=tesseract_cmd)


def _first_labeled_value(text: str, label: str) -> str:
    pattern = _LABEL_PATTERNS[label]
    for line in text.splitlines():
        match = pattern.match(line.strip())
        if match:
            return " ".join(match.group(1).split())[:300]
    return ""


def _infer_title(text: str, source: Path) -> str:
    labeled = _first_labeled_value(text, "title")
    if labeled:
        return labeled
    for line in text.splitlines():
        candidate = " ".join(line.split()).strip(" -:|")
        if 4 <= len(candidate) <= 160 and any(
            word in candidate.lower() for word in _JOB_WORDS
        ):
            return candidate
    return " ".join(source.stem.replace("_", " ").replace("-", " ").split())[:160]


def _infer_company(text: str) -> str:
    labeled = _first_labeled_value(text, "company")
    return labeled or "Unknown employer"


def _infer_location(text: str) -> str:
    return _first_labeled_value(text, "location")


def _first_url(text: str) -> str:
    match = _URL_PATTERN.search(text)
    return match.group(0).rstrip(".,);]") if match else ""


def _copy_evidence(source: Path, evidence_root: Path, digest: str) -> tuple[Path, bool]:
    destination_dir = evidence_root / digest[:2]
    destination_dir.mkdir(parents=True, exist_ok=True)
    destination = destination_dir / f"{digest}{source.suffix.lower()}"
    if destination.exists():
        return destination, False
    temporary = destination.with_suffix(destination.suffix + ".tmp")
    shutil.copyfile(source, temporary)
    temporary.replace(destination)
    return destination, True


def import_private_file(
    db_path: str | Path,
    input_path: str | Path,
    evidence_root: str | Path,
    *,
    title: str = "",
    company: str = "",
    location: str = "",
    apply_url: str = "",
    ocr: bool = False,
    tesseract_cmd: str | None = None,
    max_bytes: int = 25_000_000,
) -> PrivateImportResult:
    source = Path(input_path).resolve()
    init_private_import_tables(db_path)
    raw = source.read_bytes()
    digest = hashlib.sha256(raw).hexdigest()
    with connect(db_path) as connection:
        existing = connection.execute(
            "SELECT job_key, stored_path, review_required, warnings "
            "FROM private_imports WHERE sha256=? AND status='complete'",
            (digest,),
        ).fetchone()
    if existing:
        return PrivateImportResult(
            input_path=str(source),
            status="duplicate",
            job_key=str(existing["job_key"]),
            sha256=digest,
            stored_path=str(existing["stored_path"]),
            review_required=bool(existing["review_required"]),
            warnings=tuple(filter(None, str(existing["warnings"]).split(" | "))),
        )

    extracted = extract_document(
        source,
        ocr=ocr,
        tesseract_cmd=tesseract_cmd,
        max_bytes=max_bytes,
    )
    inferred_title = " ".join(title.split()) or _infer_title(extracted.text, source)
    inferred_company = " ".join(company.split()) or _infer_company(extracted.text)
    inferred_location = " ".join(location.split()) or _infer_location(extracted.text)
    inferred_apply_url = apply_url.strip() or _first_url(extracted.text)
    review_required = bool(extracted.warnings) or not title.strip() or not company.strip()
    warnings = list(extracted.warnings)
    if not title.strip():
        warnings.append("Job title was inferred from document text or filename")
    if not company.strip():
        warnings.append("Employer was inferred or marked Unknown employer")
    if not extracted.text.strip():
        raise ValueError("no text could be extracted from the input file")

    evidence_path, copied = _copy_evidence(source, Path(evidence_root), digest)
    import_id = uuid.uuid4().hex
    with connect(db_path) as connection:
        connection.execute(
            """
            INSERT INTO private_imports (
                import_id, sha256, original_path, original_name, stored_path,
                mime_type, size_bytes, extraction_method, extracted_chars,
                review_required, warnings, job_key, status, imported_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, '', 'staged', ?)
            """,
            (
                import_id,
                digest,
                str(source),
                source.name,
                str(evidence_path),
                extracted.mime_type,
                len(raw),
                extracted.method,
                len(extracted.text),
                int(review_required),
                " | ".join(warnings)[:4000],
                utc_now_iso(),
            ),
        )

    try:
        job = create_manual_job(
            title=inferred_title or "Unidentified vacancy",
            company=inferred_company,
            text=extracted.text,
            location=inferred_location,
            apply_url=inferred_apply_url,
            source_name=f"private:{extracted.method}",
        )
        if review_required:
            job.application_status = "review_required"
        created = upsert_job(db_path, job)
        with connect(db_path) as connection:
            connection.execute(
                "UPDATE private_imports SET job_key=?, status='complete' WHERE import_id=?",
                (job.job_key, import_id),
            )
    except Exception:
        with connect(db_path) as connection:
            connection.execute("DELETE FROM private_imports WHERE import_id=?", (import_id,))
        if copied:
            evidence_path.unlink(missing_ok=True)
        raise

    return PrivateImportResult(
        input_path=str(source),
        status="created" if created else "updated",
        job_key=job.job_key,
        created=created,
        sha256=digest,
        stored_path=str(evidence_path),
        review_required=review_required,
        warnings=tuple(warnings),
    )


def _candidate_files(folder: Path, recursive: bool) -> Iterable[Path]:
    iterator = folder.rglob("*") if recursive else folder.glob("*")
    for path in sorted(iterator):
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS:
            yield path


def import_private_folder(
    db_path: str | Path,
    folder: str | Path,
    evidence_root: str | Path,
    *,
    recursive: bool = False,
    ocr: bool = False,
    tesseract_cmd: str | None = None,
    max_files: int = 500,
    max_bytes: int = 25_000_000,
) -> FolderImportSummary:
    source_folder = Path(folder)
    if not source_folder.is_dir():
        raise NotADirectoryError(f"input folder does not exist: {source_folder}")
    files = list(_candidate_files(source_folder, recursive))
    if len(files) > max_files:
        raise ValueError(f"folder contains {len(files)} supported files; limit is {max_files}")

    summary = FolderImportSummary()
    for path in files:
        summary.attempted += 1
        try:
            result = import_private_file(
                db_path,
                path,
                evidence_root,
                ocr=ocr,
                tesseract_cmd=tesseract_cmd,
                max_bytes=max_bytes,
            )
        except Exception as exc:
            result = PrivateImportResult(
                input_path=str(path),
                status="failed",
                error_message=f"{type(exc).__name__}: {exc}",
            )
        summary.results.append(result)
        if result.status == "created":
            summary.created += 1
        elif result.status == "updated":
            summary.updated += 1
        elif result.status == "duplicate":
            summary.duplicates += 1
        else:
            summary.failed += 1
        if result.review_required:
            summary.review_required += 1
    return summary


def verify_private_imports(db_path: str | Path) -> list[str]:
    init_private_import_tables(db_path)
    errors: list[str] = []
    with connect(db_path) as connection:
        rows = connection.execute(
            "SELECT import_id, sha256, stored_path, job_key, status FROM private_imports"
        ).fetchall()
        job_keys = {
            str(row["job_key"])
            for row in connection.execute("SELECT job_key FROM jobs").fetchall()
        }
    for row in rows:
        import_id = str(row["import_id"])
        if row["status"] != "complete":
            errors.append(f"Private import is not complete: {import_id}")
            continue
        evidence_path = Path(row["stored_path"])
        if not evidence_path.exists():
            errors.append(f"Missing private import evidence: {evidence_path}")
            continue
        digest = hashlib.sha256(evidence_path.read_bytes()).hexdigest()
        if digest != row["sha256"]:
            errors.append(f"Private import hash mismatch: {evidence_path}")
        if row["job_key"] not in job_keys:
            errors.append(f"Private import has no matching job: {import_id}")
    return errors
