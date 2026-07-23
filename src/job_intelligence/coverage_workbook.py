from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font

from .coverage_overrides import apply_worldwide_company_overrides
from .coverage_proof import apply_coverage_proof
from .coverage_registry import coverage_frames
from .database import connect, fetch_jobs, init_database

REQUIRED_COVERAGE_SHEETS = (
    "Company_Registry",
    "Country_Coverage",
    "Missing_Companies",
    "Employer_Source_Status",
    "ATS_Coverage",
    "Portal_Alert_Coverage",
    "Recruiter_Source_Status",
    "PSU_Notices",
    "Source_Discovery",
    "Stale_Sources",
    "Blocked_Sources",
    "Coverage_Gaps",
    "Company_Aliases",
    "Country_Company_Matrix",
)
_FORMULA_PREFIXES = ("=", "+", "-", "@")


def _safe_value(value: object) -> object:
    if not isinstance(value, str):
        return value
    inspected = value.lstrip(" \t\r\n")
    return "'" + value if inspected.startswith(_FORMULA_PREFIXES) else value


def _safe_frame(frame: pd.DataFrame) -> pd.DataFrame:
    safe = frame.copy()
    for column in safe.columns:
        safe[column] = safe[column].map(_safe_value)
    return safe


def _format_coverage_sheets(path: Path) -> None:
    workbook = load_workbook(path)
    for sheet_name in REQUIRED_COVERAGE_SHEETS:
        if sheet_name not in workbook.sheetnames:
            continue
        sheet = workbook[sheet_name]
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for column_cells in sheet.columns:
            values = [str(cell.value or "") for cell in list(column_cells)[:200]]
            maximum = max((len(value) for value in values), default=0)
            sheet.column_dimensions[column_cells[0].column_letter].width = min(
                max(maximum + 2, 10),
                60,
            )
    workbook.save(path)


def append_coverage_sheets(
    db_path: str | Path,
    workbook_path: str | Path,
    *,
    config_dir: str | Path = "config",
) -> dict[str, pd.DataFrame]:
    database = Path(db_path)
    workbook = Path(workbook_path)
    if not workbook.exists() or workbook.stat().st_size == 0:
        raise FileNotFoundError(f"workbook does not exist or is empty: {workbook}")
    init_database(database)
    jobs = pd.DataFrame(fetch_jobs(database))
    with connect(database) as connection:
        source_health = pd.read_sql_query(
            "SELECT * FROM source_health ORDER BY source_id",
            connection,
        )
    frames = coverage_frames(jobs, source_health, config_dir=config_dir)
    frames = apply_worldwide_company_overrides(
        frames,
        source_health,
        config_dir=config_dir,
    )
    frames = apply_coverage_proof(frames)
    with pd.ExcelWriter(
        workbook,
        engine="openpyxl",
        mode="a",
        if_sheet_exists="replace",
    ) as writer:
        for sheet_name, frame in frames.items():
            _safe_frame(frame).to_excel(writer, sheet_name=sheet_name, index=False)
    _format_coverage_sheets(workbook)
    errors = verify_coverage_workbook(workbook)
    if errors:
        raise ValueError(" | ".join(errors))
    return frames


def verify_coverage_workbook(path: str | Path) -> list[str]:
    workbook_path = Path(path)
    if not workbook_path.exists():
        return [f"coverage workbook does not exist: {workbook_path}"]
    workbook = load_workbook(workbook_path, read_only=True)
    errors: list[str] = []
    missing = [
        sheet_name
        for sheet_name in REQUIRED_COVERAGE_SHEETS
        if sheet_name not in workbook.sheetnames
    ]
    if missing:
        errors.append(f"missing coverage sheets: {', '.join(missing)}")
    if "Company_Registry" in workbook.sheetnames:
        sheet = workbook["Company_Registry"]
        headers = {
            str(cell.value or "").strip()
            for cell in next(sheet.iter_rows(min_row=1, max_row=1))
        }
        required = {
            "company_id",
            "company",
            "country",
            "official_careers_url",
            "ats_family",
            "source_mode",
            "status",
            "coverage_ready",
            "coverage_proven",
            "proof_state",
            "recommended_next_action",
        }
        absent = sorted(required - headers)
        if absent:
            errors.append(f"Company_Registry missing columns: {', '.join(absent)}")
    return errors


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Append verified worldwide coverage sheets to an existing workbook."
    )
    parser.add_argument("--db", default="data/database/jobs.db")
    parser.add_argument("--workbook", default="data/exports/Piping_E3D_Jobs.xlsx")
    parser.add_argument("--config-dir", default="config")
    return parser


def main() -> int:
    args = build_parser().parse_args()
    frames = append_coverage_sheets(
        args.db,
        args.workbook,
        config_dir=args.config_dir,
    )
    registry = frames["Company_Registry"]
    proven = int(registry["coverage_proven"].sum()) if not registry.empty else 0
    ready = int(registry["coverage_ready"].sum()) if not registry.empty else 0
    print(
        f"PASS: coverage workbook verified | companies={len(registry)} | "
        f"configured={ready} | proven={proven} | workbook={args.workbook}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
