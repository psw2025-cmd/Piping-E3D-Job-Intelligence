from openpyxl import load_workbook

from job_intelligence.database import fetch_jobs, upsert_job
from job_intelligence.excel_export import REQUIRED_SHEETS, export_excel, verify_excel
from job_intelligence.manual_import import create_manual_job


def test_import_deduplicate_and_export(tmp_path) -> None:
    db_path = tmp_path / "jobs.db"
    output_path = tmp_path / "Piping_E3D_Jobs.xlsx"
    job = create_manual_job(
        title="Senior E3D Piping Designer",
        company="Example EPC",
        text="AVEVA E3D piping layout refinery vacancy. Contact jobs@example.com",
        location="Mumbai",
        apply_url="https://example.com/jobs/123?utm_source=alert",
    )

    assert upsert_job(db_path, job) is True
    assert upsert_job(db_path, job) is False
    assert len(fetch_jobs(db_path)) == 1

    export_excel(db_path, output_path)
    assert verify_excel(output_path) == []

    workbook = load_workbook(output_path, read_only=True)
    assert set(REQUIRED_SHEETS).issubset(workbook.sheetnames)
