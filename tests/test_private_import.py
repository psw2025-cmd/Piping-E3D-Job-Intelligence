from __future__ import annotations

from email.message import EmailMessage
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfWriter

from job_intelligence.database import connect, fetch_jobs
from job_intelligence.excel_export import export_excel
from job_intelligence.private_import import (
    extract_document,
    import_private_file,
    import_private_folder,
    verify_private_imports,
)


def test_text_import_is_hashed_deduplicated_and_verified(tmp_path: Path) -> None:
    source = tmp_path / "vacancy.txt"
    source.write_text(
        "Job Title: Senior E3D Piping Designer\n"
        "Company: Example EPC\n"
        "Location: Mumbai\n"
        "AVEVA E3D refinery piping layout role.\n"
        "Apply: https://example.com/jobs/123",
        encoding="utf-8",
    )
    db_path = tmp_path / "jobs.db"
    evidence = tmp_path / "private-evidence"

    first = import_private_file(
        db_path,
        source,
        evidence,
        title="Senior E3D Piping Designer",
        company="Example EPC",
        location="Mumbai",
    )
    second = import_private_file(db_path, source, evidence)

    assert first.status == "created"
    assert first.review_required is False
    assert Path(first.stored_path).exists()
    assert second.status == "duplicate"
    assert second.job_key == first.job_key
    assert len(fetch_jobs(db_path)) == 1
    assert verify_private_imports(db_path) == []
    with connect(db_path) as connection:
        assert connection.execute("SELECT COUNT(*) FROM private_imports").fetchone()[0] == 1


def test_docx_import_infers_fields_and_enters_manual_review(tmp_path: Path) -> None:
    source = tmp_path / "lead_piping_role.docx"
    document = Document()
    document.add_paragraph("Job Title: Lead Piping Engineer")
    document.add_paragraph("Company: Global EPC")
    document.add_paragraph("Location: Abu Dhabi")
    document.add_paragraph("Offshore AVEVA E3D and PDMS piping layout work.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Experience"
    table.cell(0, 1).text = "15 years"
    document.save(source)

    db_path = tmp_path / "jobs.db"
    result = import_private_file(db_path, source, tmp_path / "evidence")

    assert result.review_required is True
    job = fetch_jobs(db_path)[0]
    assert job["title"] == "Lead Piping Engineer"
    assert job["company"] == "Global EPC"
    assert job["location"] == "Abu Dhabi"
    assert job["application_status"] == "review_required"
    assert "15 years" in job["description"]

    workbook_path = tmp_path / "jobs.xlsx"
    export_excel(db_path, workbook_path)
    workbook = load_workbook(workbook_path, read_only=True)
    assert workbook["Manual_Review"].max_row == 2
    assert workbook["Private_Imports"].max_row == 2


def test_email_import_extracts_headers_body_url_and_public_email(tmp_path: Path) -> None:
    message = EmailMessage()
    message["Subject"] = "Senior Piping Engineer vacancy"
    message["From"] = "Recruiter <jobs@example.com>"
    message["To"] = "candidate@example.net"
    message.set_content(
        "Company: Example Engineering\n"
        "Location: Doha\n"
        "Senior Piping Engineer with offshore PDMS experience.\n"
        "https://example.com/apply/55"
    )
    source = tmp_path / "alert.eml"
    source.write_bytes(message.as_bytes())

    result = import_private_file(
        tmp_path / "jobs.db",
        source,
        tmp_path / "evidence",
    )

    assert result.status == "created"
    job = fetch_jobs(tmp_path / "jobs.db")[0]
    assert job["company"] == "Example Engineering"
    assert job["location"] == "Doha"
    assert job["apply_url"] == "https://example.com/apply/55"
    assert job["recruiter_email"] == "jobs@example.com"
    assert job["contact_confidence"] == "PUBLIC_UNVERIFIED"


def test_blank_pdf_is_flagged_as_needing_ocr(tmp_path: Path) -> None:
    source = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with source.open("wb") as handle:
        writer.write(handle)

    extracted = extract_document(source)

    assert extracted.text == ""
    assert any("OCR" in warning for warning in extracted.warnings)
    with pytest.raises(ValueError, match="no text"):
        import_private_file(tmp_path / "jobs.db", source, tmp_path / "evidence")


def test_image_ocr_requires_flag_and_marks_review(tmp_path: Path, monkeypatch) -> None:
    source = tmp_path / "vacancy.png"
    Image.new("RGB", (20, 20), "white").save(source)

    with pytest.raises(ValueError, match="requires --ocr"):
        extract_document(source)

    import pytesseract

    monkeypatch.setattr(
        pytesseract,
        "image_to_string",
        lambda _image: (
            "Job Title: Piping Layout Engineer\n"
            "Company: OCR EPC\n"
            "Location: Pune\n"
            "AVEVA E3D role"
        ),
    )
    result = import_private_file(
        tmp_path / "jobs.db",
        source,
        tmp_path / "evidence",
        ocr=True,
    )

    assert result.status == "created"
    assert result.review_required is True
    assert any("OCR" in warning for warning in result.warnings)


def test_folder_import_isolates_failures(tmp_path: Path) -> None:
    folder = tmp_path / "incoming"
    folder.mkdir()
    (folder / "role.txt").write_text(
        "Job Title: Piping Engineer\nCompany: Example EPC\nPDMS role",
        encoding="utf-8",
    )
    Image.new("RGB", (10, 10), "white").save(folder / "scan.png")
    (folder / "ignored.exe").write_bytes(b"not imported")

    summary = import_private_folder(
        tmp_path / "jobs.db",
        folder,
        tmp_path / "evidence",
    )

    assert summary.attempted == 2
    assert summary.created == 1
    assert summary.failed == 1
    assert len(fetch_jobs(tmp_path / "jobs.db")) == 1


def test_private_import_verification_detects_tampering(tmp_path: Path) -> None:
    source = tmp_path / "vacancy.txt"
    source.write_text(
        "Job Title: Piping Engineer\nCompany: Example EPC\nPDMS role",
        encoding="utf-8",
    )
    db_path = tmp_path / "jobs.db"
    result = import_private_file(db_path, source, tmp_path / "evidence")
    Path(result.stored_path).write_text("tampered", encoding="utf-8")

    errors = verify_private_imports(db_path)

    assert any("hash mismatch" in error for error in errors)
