from pathlib import Path

from openpyxl import load_workbook

from job_intelligence.database import upsert_job
from job_intelligence.excel_export import export_excel, verify_excel
from job_intelligence.models import JobRecord
from job_intelligence.profile_enrichment import (
    fetch_enriched_jobs,
    load_taxonomy,
    record_job_observations,
    refresh_job_enrichment,
    verify_enrichment,
)


def _config_dir() -> Path:
    return Path(__file__).resolve().parents[1] / "config"


def test_worldwide_profile_enrichment_and_cross_source_duplicate(tmp_path: Path) -> None:
    db_path = tmp_path / "jobs.db"
    job = JobRecord(
        title="Lead Piping Design Engineer",
        company="Example EPC",
        location="Bangalore, India",
        description=(
            "Oil & gas brownfield role using AVEVA E3D, SP3D and Navisworks. "
            "Minimum 12 years experience. Application deadline: 31/12/2026. "
            "Hybrid working with site engineering support."
        ),
        apply_url="https://example.com/jobs/lead-piping-1",
        source_url="https://example.com/jobs/lead-piping-1",
        source_name="Official employer",
        published_at="2026-07-22",
    )
    assert upsert_job(db_path, job)
    record_job_observations(
        db_path,
        [job],
        source_id="official_epc",
        source_type="workday",
    )
    record_job_observations(
        db_path,
        [job],
        source_id="recruiter_alert",
        source_type="gmail_alert",
    )

    refresh_job_enrichment(db_path, _config_dir())
    rows = fetch_enriched_jobs(db_path)

    assert len(rows) == 1
    row = rows[0]
    assert row["normalized_role"] == "Lead Piping Engineer"
    assert row["city"] == "Bengaluru"
    assert row["country"] == "India"
    assert row["closing_date"] == "2026-12-31"
    assert row["experience_required"] == "12+ years"
    assert "AVEVA E3D" in row["software"]
    assert "Hexagon Smart 3D" in row["software"]
    assert "Navisworks" in row["software"]
    assert "Oil and Gas" in row["sector"]
    assert row["employment_mode"] == "Hybrid"
    assert row["duplicate_status"] == "cross_source"
    assert row["duplicate_source_count"] == 2
    assert row["canonical_job_key"] == job.job_key
    assert verify_enrichment(db_path) == []


def test_complete_role_and_location_matrices_are_configured() -> None:
    taxonomy = load_taxonomy(_config_dir())
    expected_roles = {
        "Lead Piping Engineer",
        "Senior Piping Engineer",
        "Piping Design Engineer",
        "Piping Layout Engineer",
        "Piping Design Checker",
        "Site Piping Engineer",
        "Offshore Piping Engineer",
        "Senior E3D Piping Designer",
        "Senior PDMS Designer",
        "SP3D Piping Designer",
        "E3D Administrator",
        "E3D Coordinator",
        "Plant Design Engineer",
        "Piping CAD Lead",
    }
    assert expected_roles <= set(taxonomy.role_families)

    expected_countries = {
        "India",
        "United Arab Emirates",
        "Saudi Arabia",
        "Qatar",
        "Oman",
        "Kuwait",
        "Bahrain",
        "United Kingdom",
        "Canada",
        "Australia",
        "Singapore",
        "Malaysia",
        "Japan",
        "South Korea",
        "United States",
    }
    assert expected_countries <= set(taxonomy.country_aliases.values())


def test_worldwide_excel_contains_normalized_fields_and_summary(tmp_path: Path) -> None:
    db_path = tmp_path / "jobs.db"
    output = tmp_path / "worldwide.xlsx"
    job = JobRecord(
        title="Senior E3D Piping Designer",
        company="Example EPC",
        location="Abu Dhabi, UAE",
        description="Offshore LNG piping layout using AVEVA E3D and AutoCAD.",
        apply_url="https://example.com/jobs/e3d-1",
        source_name="Official employer",
    )
    upsert_job(db_path, job)
    record_job_observations(
        db_path,
        [job],
        source_id="official_epc",
        source_type="oracle_hcm",
    )

    export_excel(db_path, output)

    assert verify_excel(output) == []
    workbook = load_workbook(output, read_only=True)
    assert "Worldwide_Active" in workbook.sheetnames
    assert "Cross_Source_Duplicates" in workbook.sheetnames
    assert "Global_Summary" in workbook.sheetnames
    headers = [cell.value for cell in next(workbook["Worldwide_Active"].iter_rows())]
    for required in (
        "normalized_role",
        "city",
        "country",
        "source_type",
        "closing_date",
        "experience_required",
        "software",
        "sector",
        "employment_mode",
        "duplicate_status",
        "duplicate_source_count",
        "canonical_job_key",
    ):
        assert required in headers
