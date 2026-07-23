from pathlib import Path

from openpyxl import load_workbook

from job_intelligence.database import upsert_jobs
from job_intelligence.excel_export import export_excel, verify_excel
from job_intelligence.models import JobRecord


def _headers(sheet) -> dict[str, int]:
    return {str(cell.value): cell.column for cell in sheet[1]}


def test_worldwide_workbook_groups_cross_source_duplicates(tmp_path: Path) -> None:
    db_path = tmp_path / "jobs.db"
    jobs = [
        JobRecord(
            title="Senior E3D Piping Designer",
            company="Global EPC",
            location="Abu Dhabi, UAE",
            description="AVEVA E3D offshore piping layout role. Apply by 31 July 2026.",
            apply_url="https://company.example/jobs/123",
            source_name="Official employer",
            published_at="2026-07-20",
            match_score=90,
            priority="critical",
        ),
        JobRecord(
            title="Senior E3D Piping Designer",
            company="Global EPC",
            location="Abu Dhabi, UAE",
            description="AVEVA E3D offshore piping layout role from recruiter alert.",
            apply_url="https://recruiter.example/jobs/abc",
            source_name="Airswift",
            published_at="2026-07-20",
            match_score=80,
            priority="high",
        ),
    ]
    upsert_jobs(db_path, jobs)
    output = tmp_path / "worldwide.xlsx"

    export_excel(db_path, output)

    assert verify_excel(output) == []
    workbook = load_workbook(output, read_only=True)
    dedup = workbook["Worldwide_Dedup"]
    variants = workbook["Duplicate_Variants"]
    assert dedup.max_row == 2
    assert variants.max_row == 3

    headers = _headers(dedup)
    row = dedup[2]
    assert row[headers["source_count"] - 1].value == 2
    assert row[headers["country"] - 1].value == "United Arab Emirates"
    assert row[headers["sector"] - 1].value == "Offshore"
    assert "AVEVA E3D" in str(row[headers["software"] - 1].value)
    assert row[headers["normalized_role"] - 1].value == "Senior E3D/PDMS/SP3D Designer"
    assert "Official employer" in str(row[headers["duplicate_sources"] - 1].value)
    assert "Airswift" in str(row[headers["duplicate_sources"] - 1].value)


def test_empty_worldwide_workbook_keeps_full_contract(tmp_path: Path) -> None:
    output = tmp_path / "empty.xlsx"

    export_excel(tmp_path / "empty.db", output)

    assert verify_excel(output) == []
    workbook = load_workbook(output, read_only=True)
    headers = _headers(workbook["Worldwide_Dedup"])
    for required in (
        "normalized_role",
        "country",
        "sector",
        "software",
        "closing_date",
        "source_category",
        "duplicate_group",
        "source_count",
        "duplicate_sources",
        "all_apply_urls",
    ):
        assert required in headers
