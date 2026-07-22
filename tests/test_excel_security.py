from pathlib import Path

from openpyxl import load_workbook

from job_intelligence.database import upsert_job
from job_intelligence.excel_export import export_excel
from job_intelligence.models import JobRecord


def test_external_text_cannot_become_excel_formula(tmp_path: Path) -> None:
    db_path = tmp_path / "jobs.db"
    output = tmp_path / "jobs.xlsx"
    upsert_job(
        db_path,
        JobRecord(
            title='=HYPERLINK("https://malicious.example", "Open")',
            company="Example EPC",
            description="+SUM(1,1)",
        ),
    )

    export_excel(db_path, output)

    workbook = load_workbook(output, data_only=False)
    sheet = workbook["All_Active"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    title_cell = sheet.cell(row=2, column=headers["title"])
    description_cell = sheet.cell(row=2, column=headers["description"])

    assert title_cell.data_type == "s"
    assert str(title_cell.value).startswith("'=")
    assert description_cell.data_type == "s"
    assert str(description_cell.value).startswith("'+")
