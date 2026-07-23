from __future__ import annotations

import base64
from email.message import EmailMessage
from pathlib import Path

from openpyxl import load_workbook

from job_intelligence.database import connect, fetch_jobs
from job_intelligence.excel_export import export_excel, verify_excel
from job_intelligence.gmail_alerts import (
    extract_alert_candidates,
    import_gmail_raw_message,
    import_gmail_service,
    verify_gmail_imports,
)


def _alert_email() -> bytes:
    message = EmailMessage()
    message["Subject"] = "Two piping jobs for you"
    message["From"] = "LinkedIn Jobs <alerts@linkedin.com>"
    message["To"] = "candidate@example.com"
    message["Date"] = "Thu, 23 Jul 2026 09:00:00 +0530"
    message.set_content("Your HTML job alert is attached below.")
    message.add_alternative(
        """
        <html><body>
          <article>
            <p>Company: Example EPC</p>
            <p>Location: Mumbai, India</p>
            <a href="https://example.com/jobs/lead-piping-123">
              Lead Piping Engineer
            </a>
          </article>
          <article>
            <p>Company: Gulf Engineering</p>
            <p>Location: Abu Dhabi, UAE</p>
            <a href="https://gulf.example/jobs/e3d-designer-55">
              Senior E3D Piping Designer
            </a>
          </article>
          <a href="https://linkedin.com/settings/email">Email settings</a>
          <a href="https://linkedin.com/unsubscribe">Unsubscribe</a>
        </body></html>
        """,
        subtype="html",
    )
    return message.as_bytes()


def test_html_alert_extracts_multiple_jobs_and_ignores_account_links() -> None:
    _, candidates = extract_alert_candidates(_alert_email())

    assert len(candidates) == 2
    assert {candidate.title for candidate in candidates} == {
        "Lead Piping Engineer",
        "Senior E3D Piping Designer",
    }
    assert {candidate.company for candidate in candidates} == {
        "Example EPC",
        "Gulf Engineering",
    }
    assert all("unsubscribe" not in candidate.url for candidate in candidates)


def test_gmail_message_import_is_incremental_hashed_and_verified(tmp_path: Path) -> None:
    db_path = tmp_path / "jobs.db"
    evidence = tmp_path / "gmail-evidence"
    raw = _alert_email()

    first = import_gmail_raw_message(
        db_path,
        message_id="gmail-1",
        thread_id="thread-1",
        raw_message=raw,
        evidence_root=evidence,
        query_text="piping jobs",
    )
    second = import_gmail_raw_message(
        db_path,
        message_id="gmail-1",
        thread_id="thread-1",
        raw_message=raw,
        evidence_root=evidence,
        query_text="piping jobs",
    )

    assert first.status == "complete"
    assert first.created == 2
    assert len(first.job_keys) == 2
    assert Path(first.evidence_path).is_file()
    assert second.status == "duplicate"
    assert len(fetch_jobs(db_path)) == 2
    assert verify_gmail_imports(db_path) == []

    with connect(db_path) as connection:
        assert connection.execute(
            "SELECT COUNT(*) FROM gmail_alert_messages"
        ).fetchone()[0] == 1
        assert connection.execute(
            "SELECT COUNT(*) FROM gmail_alert_jobs"
        ).fetchone()[0] == 2


def test_unknown_employer_alert_is_routed_to_manual_review(tmp_path: Path) -> None:
    message = EmailMessage()
    message["Subject"] = "Piping opportunity"
    message["From"] = "Job Alerts <alerts@example.net>"
    message["To"] = "candidate@example.com"
    message.set_content(
        "Senior Piping Engineer\n"
        "Location: Doha, Qatar\n"
        "https://jobs.example.net/job/100"
    )

    result = import_gmail_raw_message(
        tmp_path / "jobs.db",
        message_id="gmail-unknown",
        raw_message=message.as_bytes(),
        evidence_root=tmp_path / "evidence",
    )

    assert result.created == 1
    job = fetch_jobs(tmp_path / "jobs.db")[0]
    assert job["company"] == "Unknown employer"
    assert job["application_status"] == "review_required"
    assert job["contact_confidence"] == ""


def test_deleted_linked_job_is_recreated_on_reimport(tmp_path: Path) -> None:
    db_path = tmp_path / "jobs.db"
    raw = _alert_email()
    first = import_gmail_raw_message(
        db_path,
        message_id="gmail-repair",
        raw_message=raw,
        evidence_root=tmp_path / "evidence",
    )
    with connect(db_path) as connection:
        connection.execute("DELETE FROM job_identity_aliases")
        connection.execute("DELETE FROM jobs WHERE job_key=?", (first.job_keys[0],))

    repaired = import_gmail_raw_message(
        db_path,
        message_id="gmail-repair",
        raw_message=raw,
        evidence_root=tmp_path / "evidence",
    )

    assert repaired.status == "complete"
    assert len(fetch_jobs(db_path)) == 2
    assert verify_gmail_imports(db_path) == []


class _Execute:
    def __init__(self, payload: dict[str, object]) -> None:
        self.payload = payload

    def execute(self) -> dict[str, object]:
        return self.payload


class _Messages:
    def __init__(self, raw_messages: dict[str, bytes]) -> None:
        self.raw_messages = raw_messages

    def list(self, **_kwargs) -> _Execute:
        return _Execute({"messages": [{"id": key} for key in self.raw_messages]})

    def get(self, *, id: str, **_kwargs) -> _Execute:
        encoded = base64.urlsafe_b64encode(self.raw_messages[id]).decode("ascii").rstrip("=")
        return _Execute(
            {
                "id": id,
                "threadId": f"thread-{id}",
                "internalDate": "1784777400000",
                "raw": encoded,
            }
        )


class _Users:
    def __init__(self, raw_messages: dict[str, bytes]) -> None:
        self._messages = _Messages(raw_messages)

    def messages(self) -> _Messages:
        return self._messages


class FakeGmailService:
    def __init__(self, raw_messages: dict[str, bytes]) -> None:
        self._users = _Users(raw_messages)

    def users(self) -> _Users:
        return self._users


def test_gmail_service_import_and_workbook_proof(tmp_path: Path) -> None:
    db_path = tmp_path / "jobs.db"
    service = FakeGmailService({"one": _alert_email()})

    summary = import_gmail_service(
        service,
        db_path,
        evidence_root=tmp_path / "gmail-evidence",
        query="newer_than:30d piping",
        max_messages=10,
    )

    assert summary.attempted == 1
    assert summary.processed == 1
    assert summary.jobs_created == 2
    workbook_path = tmp_path / "jobs.xlsx"
    export_excel(db_path, workbook_path)
    assert verify_excel(workbook_path) == []
    workbook = load_workbook(workbook_path, read_only=True)
    assert workbook["Gmail_Alerts"].max_row == 2
    assert workbook["Gmail_Alert_Jobs"].max_row == 3
    assert workbook["Daily_Summary"].max_row > 2
