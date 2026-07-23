from __future__ import annotations

import sqlite3
from pathlib import Path

from openpyxl import load_workbook

from job_intelligence.database import fetch_jobs, init_database, upsert_job
from job_intelligence.excel_export import REQUIRED_JOB_COLUMNS, export_excel, verify_excel
from job_intelligence.models import JobRecord
from job_intelligence.normalization import enrich_job, load_profile_taxonomy
from job_intelligence.scoring import score_job


CONFIG_DIR = Path(__file__).resolve().parents[1] / "config"


def test_worldwide_role_location_and_vacancy_fields_are_normalized() -> None:
    taxonomy = load_profile_taxonomy(CONFIG_DIR)
    job = JobRecord(
        title="Lead Piping Design Engineer",
        company="Global EPC",
        location="Abu Dhabi, UAE",
        description=(
            "Contract LNG project requiring AVEVA E3D, PDMS, plant layout and "
            "offshore model review. Closing date: 31 August 2026"
        ),
    )

    enrich_job(job, config_dir=CONFIG_DIR, taxonomy=taxonomy)

    assert job.normalized_role == "Lead Piping Design Engineer"
    assert job.city == "Abu Dhabi"
    assert job.country == "United Arab Emirates"
    assert job.software_text == "AVEVA E3D; PDMS"
    assert job.sector == "Offshore" or job.sector == "LNG"
    assert job.employment_type == "Contract"
    assert job.closing_at == "31 August 2026"
    assert job.duplicate_status == "unique"


def test_abbreviations_and_city_aliases_match_profile() -> None:
    job = JobRecord(
        title="SP3D Piping Designer",
        company="Example Engineering",
        location="Bangalore, India",
        description="SmartPlant 3D refinery piping and equipment layout position.",
    )

    result = score_job(job, config_dir=CONFIG_DIR)

    assert job.normalized_role == "SP3D Piping Designer"
    assert job.city == "Bengaluru"
    assert job.country == "India"
    assert "SP3D" in job.software_text
    assert job.sector == "Refinery"
    assert result.score >= 50
    assert any("role" in reason.lower() for reason in result.reasons)


def test_unrelated_role_is_not_normalized_to_piping() -> None:
    job = JobRecord(
        title="Accountant",
        company="Example Company",
        location="Remote in Europe",
        description="General ledger and taxation work for an engineering company.",
    )

    enrich_job(job, config_dir=CONFIG_DIR)

    assert job.normalized_role == ""
    assert job.country == "Europe"


def test_existing_database_migrates_worldwide_fields(tmp_path: Path) -> None:
    db_path = tmp_path / "legacy.db"
    connection = sqlite3.connect(db_path)
    connection.execute(
        """
        CREATE TABLE jobs (
            job_key TEXT PRIMARY KEY,
            title TEXT NOT NULL,
            company TEXT NOT NULL,
            location TEXT NOT NULL DEFAULT '',
            description TEXT NOT NULL DEFAULT '',
            apply_url TEXT NOT NULL DEFAULT '',
            source_url TEXT NOT NULL DEFAULT '',
            source_name TEXT NOT NULL DEFAULT '',
            published_at TEXT NOT NULL DEFAULT '',
            found_at TEXT NOT NULL,
            last_seen_at TEXT NOT NULL,
            job_type TEXT NOT NULL DEFAULT '',
            salary_text TEXT NOT NULL DEFAULT '',
            experience_text TEXT NOT NULL DEFAULT '',
            skills_text TEXT NOT NULL DEFAULT '',
            recruiter_name TEXT NOT NULL DEFAULT '',
            recruiter_email TEXT NOT NULL DEFAULT '',
            contact_confidence TEXT NOT NULL DEFAULT '',
            match_score INTEGER NOT NULL DEFAULT 0,
            match_reasons TEXT NOT NULL DEFAULT '',
            gaps TEXT NOT NULL DEFAULT '',
            priority TEXT NOT NULL DEFAULT 'normal',
            application_status TEXT NOT NULL DEFAULT 'new'
        )
        """
    )
    connection.commit()
    connection.close()

    init_database(db_path)

    connection = sqlite3.connect(db_path)
    columns = {row[1] for row in connection.execute("PRAGMA table_info(jobs)")}
    connection.close()
    assert set(REQUIRED_JOB_COLUMNS) <= columns


def test_excel_contains_final_tracking_fields_and_sheets(tmp_path: Path) -> None:
    db_path = tmp_path / "jobs.db"
    job = JobRecord(
        title="Senior E3D Piping Designer",
        company="Example EPC",
        location="Mumbai, India",
        description="Permanent AVEVA E3D petrochemical piping layout role.",
        apply_url="https://example.com/jobs/123",
        source_url="https://example.com/careers",
        source_name="official employer",
    )
    score_job(job, config_dir=CONFIG_DIR)
    upsert_job(db_path, job)

    workbook_path = tmp_path / "Piping_E3D_Jobs.xlsx"
    export_excel(db_path, workbook_path)

    assert verify_excel(workbook_path) == []
    workbook = load_workbook(workbook_path, read_only=True)
    assert "Rejected" in workbook.sheetnames
    assert "Daily_Summary" in workbook.sheetnames
    headers = {
        str(cell.value)
        for cell in next(workbook["All_Active"].iter_rows(min_row=1, max_row=1))
    }
    assert set(REQUIRED_JOB_COLUMNS) <= headers
    stored = fetch_jobs(db_path)[0]
    assert stored["normalized_role"] == "Senior E3D Piping Designer"
    assert stored["city"] == "Mumbai"
    assert stored["country"] == "India"
