from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from job_intelligence.coverage_workbook import (
    REQUIRED_COVERAGE_SHEETS,
    append_coverage_sheets,
    verify_coverage_workbook,
)
from job_intelligence.database import upsert_job
from job_intelligence.excel_export import export_excel
from job_intelligence.models import JobRecord


def repository_root() -> Path:
    return Path(__file__).resolve().parents[1]


def test_coverage_sheets_are_appended_and_verified(tmp_path: Path) -> None:
    db_path = tmp_path / "jobs.db"
    workbook_path = tmp_path / "jobs.xlsx"
    upsert_job(
        db_path,
        JobRecord(
            title="Senior Piping Engineer",
            company="McDermott",
            country="India",
            location="Chennai, India",
            description="AVEVA E3D piping layout role",
            apply_url="https://example.com/job/1",
            source_url="https://example.com/job/1",
            source_name="fixture",
        ),
    )
    export_excel(db_path, workbook_path)

    frames = append_coverage_sheets(
        db_path,
        workbook_path,
        config_dir=repository_root() / "config",
    )

    assert verify_coverage_workbook(workbook_path) == []
    workbook = load_workbook(workbook_path, read_only=True)
    assert set(REQUIRED_COVERAGE_SHEETS) <= set(workbook.sheetnames)
    assert len(frames["Company_Registry"]) >= 31
    assert len(frames["Country_Coverage"]) >= 2


def test_coverage_workbook_neutralizes_formula_text(tmp_path: Path) -> None:
    db_path = tmp_path / "jobs.db"
    workbook_path = tmp_path / "jobs.xlsx"
    export_excel(db_path, workbook_path)
    append_coverage_sheets(
        db_path,
        workbook_path,
        config_dir=repository_root() / "config",
    )

    workbook = load_workbook(workbook_path, data_only=False)
    sheet = workbook["Portal_Alert_Coverage"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    query_cell = sheet.cell(row=2, column=headers["query"])
    assert query_cell.data_type == "s"
